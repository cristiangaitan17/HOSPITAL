import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from django.http import HttpResponse
from .models import Especialidad, Medico, Paciente, Cita, Medicamento, Tratamiento, Factura, Pago

def export_to_excel(queryset, model_name, filename):
    """Exporta un queryset a Excel con formato profesional"""
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = model_name
    
    # Obtener campos del modelo
    fields = [f.name for f in queryset.model._meta.get_fields() 
              if not f.auto_created and not f.is_relation]
    
    # Estilo para encabezados
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    header_alignment = Alignment(horizontal='center', vertical='center')
    
    # Encabezados
    for col, field in enumerate(fields, 1):
        cell = ws.cell(row=1, column=col, value=field)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    # Datos
    for row, obj in enumerate(queryset, 2):
        for col, field in enumerate(fields, 1):
            value = getattr(obj, field)
            if value is None:
                value = ''
            elif hasattr(value, 'strftime'):  # Fechas
                value = value.strftime('%Y-%m-%d %H:%M:%S')
            ws.cell(row=row, column=col, value=str(value))
    
    # Ajustar columnas
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 40)
        ws.column_dimensions[col_letter].width = adjusted_width
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename={filename}.xlsx'
    wb.save(response)
    return response


def export_especialidades(request):
    return export_to_excel(Especialidad.objects.all(), 'Especialidades', 'especialidades')


def export_medicos(request):
    return export_to_excel(Medico.objects.all(), 'Medicos', 'medicos')


def export_pacientes(request):
    return export_to_excel(Paciente.objects.all(), 'Pacientes', 'pacientes')


def export_citas(request):
    return export_to_excel(Cita.objects.all(), 'Citas', 'citas')


def export_medicamentos(request):
    return export_to_excel(Medicamento.objects.all(), 'Medicamentos', 'medicamentos')


def export_tratamientos(request):
    return export_to_excel(Tratamiento.objects.all(), 'Tratamientos', 'tratamientos')


def export_facturas(request):
    return export_to_excel(Factura.objects.all(), 'Facturas', 'facturas')


def export_pagos(request):
    return export_to_excel(Pago.objects.all(), 'Pagos', 'pagos')